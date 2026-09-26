"""
Servicio de Generación de Reportes Periódicos en Excel (.xlsx) para DiarioComercial.
Genera libros de cálculo multi-hoja con diseño ejecutivo para comerciantes, dueños y contadores.
"""
import io
from datetime import date, datetime
from decimal import Decimal
from typing import Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from django.db.models import Sum
from .models import Compra, Establecimiento, Producto, Retencion, Venta


def generar_excel_reporte_periodico(
    est: Establecimiento,
    desde,
    hasta,
    titulo_periodo: str = "Reporte de Operaciones"
) -> Tuple[bytes, str]:
    """
    Genera un archivo Excel (.xlsx) completo con 4 hojas:
      1. Resumen y Arqueo de Caja
      2. Detalle de Ventas
      3. Detalle de Compras y Gastos
      4. Inventario y Existencias Valorizadas
    Retorna (archivo_bytes, nombre_archivo).
    """
    wb = openpyxl.Workbook()

    # Formateo de fechas para queries
    if isinstance(desde, datetime):
        dt_desde = desde
        d_desde = desde.date()
    else:
        d_desde = desde
        dt_desde = datetime.combine(desde, datetime.min.time())

    if isinstance(hasta, datetime):
        dt_hasta = hasta
        d_hasta = hasta.date()
    else:
        d_hasta = hasta
        dt_hasta = datetime.combine(hasta, datetime.max.time())

    # Paleta de Estilos Oficial de DiarioComercial
    fill_header = PatternFill(start_color="12323A", end_color="12323A", fill_type="solid")
    fill_accent = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    fill_subtotal = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_green = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_titulo = Font(name="Calibri", size=14, bold=True, color="12323A")
    font_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_muted = Font(name="Calibri", size=9, italic=True, color="64748B")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    double_bottom_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="double", color="0F172A"),
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    fmt_moneda = "$#,##0"
    fmt_decimal = "#,##0.00"

    # Consultas filtradas estrictamente por el establecimiento
    ventas_qs = Venta.objects.filter(
        establecimiento=est,
        fecha__range=(d_desde, d_hasta),
        estado="vigente"
    ).select_related("cliente", "usuario", "actividad").order_by("-fecha_hora", "-id")

    compras_qs = Compra.objects.filter(
        establecimiento=est,
        fecha__range=(d_desde, d_hasta),
        estado="vigente"
    ).select_related("usuario").order_by("-fecha", "-id")

    retenciones_qs = Retencion.objects.filter(
        establecimiento=est,
        fecha__range=(d_desde, d_hasta),
        estado="vigente"
    ).order_by("-fecha", "-id")

    productos_qs = Producto.objects.filter(
        establecimiento=est,
        estado="activo"
    ).order_by("categoria", "nombre")

    total_ventas = ventas_qs.aggregate(t=Sum("valor"))["t"] or Decimal("0")
    total_ica = ventas_qs.aggregate(t=Sum("ica_estimado"))["t"] or Decimal("0")
    total_compras = compras_qs.aggregate(t=Sum("valor"))["t"] or Decimal("0")
    total_rets = retenciones_qs.aggregate(t=Sum("valor"))["t"] or Decimal("0")
    flujo_neto = total_ventas - total_compras
    cant_ventas = ventas_qs.count()
    ticket_promedio = (total_ventas / cant_ventas) if cant_ventas > 0 else Decimal("0")

    # =========================================================================
    # HOJA 1: RESUMEN EJECUTIVO Y ARQUEO
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Resumen y Arqueo"
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = "DIARIO COMERCIAL — REPORTE AUTOMÁTICO DE OPERACIONES"
    ws1["A1"].font = font_titulo
    ws1["A2"] = f"Comercio: {est.nombre} | NIT: {est.nit or 'No registrado'} | Municipio: {est.municipio.nombre if est.municipio else 'Tunja'}"
    ws1["A2"].font = font_bold
    ws1["A3"] = f"Periodo: {titulo_periodo} ({d_desde.strftime('%d/%m/%Y')} al {d_hasta.strftime('%d/%m/%Y')}) | Emitido: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws1["A3"].font = font_muted

    # Tabla Resumen General
    headers_kpi = ["Métrica Financiera", "Valor ($ COP)", "Detalle"]
    for col_num, h in enumerate(headers_kpi, start=1):
        cell = ws1.cell(row=5, column=col_num, value=h)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center

    kpis = [
        ("Ventas Totales Brutas", total_ventas, f"{cant_ventas} transacciones registradas"),
        ("Compras y Gastos Operativos", total_compras, f"{compras_qs.count()} gastos registrados"),
        ("Balance Neto de Caja", flujo_neto, "Ingresos brutos menos gastos"),
        ("Ticket Promedio por Venta", ticket_promedio, "Valor promedio por compra en mostrador"),
        ("Retenciones Practicadas a Favor", total_rets, f"{retenciones_qs.count()} retenciones aplicadas"),
        ("Impuesto ICA Estimado del Periodo", total_ica, "Provisión estimada para declaración municipal"),
    ]

    r = 6
    for titulo, valor, nota in kpis:
        c1 = ws1.cell(row=r, column=1, value=titulo)
        c2 = ws1.cell(row=r, column=2, value=float(valor))
        c3 = ws1.cell(row=r, column=3, value=nota)

        c1.border = thin_border
        c2.border = thin_border
        c3.border = thin_border
        c2.number_format = fmt_moneda
        c2.alignment = align_right

        if titulo == "Balance Neto de Caja":
            c1.font = font_bold
            c2.font = font_bold
            c1.fill = fill_green
            c2.fill = fill_green
        r += 1

    # Tabla Arqueo por Medio de Pago
    r += 2
    ws1.cell(row=r, column=1, value="ARQUEO POR MEDIO DE PAGO").font = font_bold
    r += 1

    headers_medios = ["Medio de Pago", "Total Recaudado ($ COP)", "Participación"]
    for col_num, h in enumerate(headers_medios, start=1):
        cell = ws1.cell(row=r, column=col_num, value=h)
        cell.fill = fill_accent
        cell.font = font_header
        cell.alignment = align_center
    r += 1

    medios_dict = {
        "efectivo": "Efectivo en Cajón",
        "nequi": "Nequi",
        "daviplata": "Daviplata",
        "transferencia": "Transferencia Bancaria",
        "bre_b": "Bre-B (Interoperable BanRep)",
    }
    for cod_medio, nombre_medio in medios_dict.items():
        v_sub = ventas_qs.filter(medio_pago=cod_medio).aggregate(t=Sum("valor"))["t"] or Decimal("0")
        pct = (v_sub / total_ventas * 100) if total_ventas > 0 else Decimal("0")

        c1 = ws1.cell(row=r, column=1, value=nombre_medio)
        c2 = ws1.cell(row=r, column=2, value=float(v_sub))
        c3 = ws1.cell(row=r, column=3, value=f"{pct:.1f}%")

        c1.border = thin_border
        c2.border = thin_border
        c3.border = thin_border
        c2.number_format = fmt_moneda
        c2.alignment = align_right
        c3.alignment = align_center
        r += 1

    # =========================================================================
    # HOJA 2: DETALLE DE VENTAS
    # =========================================================================
    ws2 = wb.create_sheet(title="Detalle Ventas")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = f"DETALLE DE VENTAS — {est.nombre}"
    ws2["A1"].font = font_titulo
    ws2["A2"] = f"Periodo: {d_desde.strftime('%d/%m/%Y')} al {d_hasta.strftime('%d/%m/%Y')}"
    ws2["A2"].font = font_muted

    headers_v = [
        "Comprobante #", "Fecha y Hora", "Concepto / Detalle", "Medio de Pago",
        "Valor Venta ($)", "ICA Estimado ($)", "Cliente / Adquirente", "Factura Electrónica / CUFE"
    ]
    for col_num, h in enumerate(headers_v, start=1):
        cell = ws2.cell(row=4, column=col_num, value=h)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center

    row_v = 5
    for v in ventas_qs:
        cli_str = v.cliente.nombre if v.cliente else "Consumidor Final"
        fe_str = f"{v.numero_factura_electronica} ({v.estado_dian})" if v.numero_factura_electronica else "Mostrador POS"

        ws2.cell(row=row_v, column=1, value=v.pk).alignment = align_center
        ws2.cell(row=row_v, column=2, value=v.fecha_hora.strftime("%d/%m/%Y %H:%M")).alignment = align_center
        ws2.cell(row=row_v, column=3, value=v.concepto)
        ws2.cell(row=row_v, column=4, value=v.get_medio_pago_display().upper()).alignment = align_center

        c_val = ws2.cell(row=row_v, column=5, value=float(v.valor))
        c_val.number_format = fmt_moneda
        c_val.alignment = align_right

        c_ica = ws2.cell(row=row_v, column=6, value=float(v.ica_estimado))
        c_ica.number_format = fmt_moneda
        c_ica.alignment = align_right

        ws2.cell(row=row_v, column=7, value=cli_str)
        ws2.cell(row=row_v, column=8, value=fe_str)

        for c in range(1, 9):
            ws2.cell(row=row_v, column=c).border = thin_border
        row_v += 1

    # Fila de totales en ventas
    ws2.cell(row=row_v, column=3, value="TOTAL VENTAS DEL PERIODO").font = font_bold
    c_tot_v = ws2.cell(row=row_v, column=5, value=float(total_ventas))
    c_tot_v.font = font_bold
    c_tot_v.number_format = fmt_moneda
    c_tot_v.border = double_bottom_border

    c_tot_ica = ws2.cell(row=row_v, column=6, value=float(total_ica))
    c_tot_ica.font = font_bold
    c_tot_ica.number_format = fmt_moneda
    c_tot_ica.border = double_bottom_border

    # =========================================================================
    # HOJA 3: COMPRAS Y GASTOS
    # =========================================================================
    ws3 = wb.create_sheet(title="Compras y Gastos")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = f"COMPRAS Y GASTOS OPERATIVOS — {est.nombre}"
    ws3["A1"].font = font_titulo
    ws3["A2"] = f"Periodo: {d_desde.strftime('%d/%m/%Y')} al {d_hasta.strftime('%d/%m/%Y')}"
    ws3["A2"].font = font_muted

    headers_c = ["ID Gasto #", "Fecha", "Proveedor / Concepto", "Valor Gasto ($)", "Registrado Por"]
    for col_num, h in enumerate(headers_c, start=1):
        cell = ws3.cell(row=4, column=col_num, value=h)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center

    row_c = 5
    for c in compras_qs:
        ws3.cell(row=row_c, column=1, value=c.pk).alignment = align_center
        ws3.cell(row=row_c, column=2, value=c.fecha.strftime("%d/%m/%Y")).alignment = align_center
        ws3.cell(row=row_c, column=3, value=c.proveedor)

        c_val = ws3.cell(row=row_c, column=4, value=float(c.valor))
        c_val.number_format = fmt_moneda
        c_val.alignment = align_right

        ws3.cell(row=row_c, column=5, value=c.usuario.username if c.usuario else "Sistema")

        for col in range(1, 6):
            ws3.cell(row=row_c, column=col).border = thin_border
        row_c += 1

    ws3.cell(row=row_c, column=3, value="TOTAL GASTOS Y COMPRAS").font = font_bold
    c_tot_c = ws3.cell(row=row_c, column=4, value=float(total_compras))
    c_tot_c.font = font_bold
    c_tot_c.number_format = fmt_moneda
    c_tot_c.border = double_bottom_border

    # =========================================================================
    # HOJA 4: INVENTARIO ACTUAL Y EXISTENCIAS
    # =========================================================================
    ws4 = wb.create_sheet(title="Inventario Actual")
    ws4.views.sheetView[0].showGridLines = True

    ws4["A1"] = f"CATÁLOGO Y EXISTENCIAS EN BODEGA — {est.nombre}"
    ws4["A1"].font = font_titulo
    ws4["A2"] = f"Corte a fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Total Artículos: {productos_qs.count()}"
    ws4["A2"].font = font_muted

    headers_i = [
        "Código / SKU", "Producto", "Categoría", "Stock Actual",
        "Unidad", "Costo Unitario ($)", "Precio Venta ($)", "Margen (%)", "Valor Total Stock ($)"
    ]
    for col_num, h in enumerate(headers_i, start=1):
        cell = ws4.cell(row=4, column=col_num, value=h)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center

    row_i = 5
    total_val_bodega = Decimal("0")

    for p in productos_qs:
        subtotal_stock = p.stock_kilos * (p.costo_unitario if p.costo_unitario > 0 else p.precio_kilo)
        total_val_bodega += subtotal_stock

        margen = Decimal("0")
        if p.costo_unitario > 0 and p.precio_kilo > p.costo_unitario:
            margen = ((p.precio_kilo - p.costo_unitario) / p.precio_kilo) * 100

        ws4.cell(row=row_i, column=1, value=p.codigo_barras or f"PRD-{p.pk:04d}").alignment = align_center
        ws4.cell(row=row_i, column=2, value=p.nombre)
        ws4.cell(row=row_i, column=3, value=p.get_categoria_display()).alignment = align_center

        c_stock = ws4.cell(row=row_i, column=4, value=float(p.stock_kilos))
        c_stock.number_format = fmt_decimal
        c_stock.alignment = align_right

        ws4.cell(row=row_i, column=5, value=p.unidad_medida).alignment = align_center

        c_costo = ws4.cell(row=row_i, column=6, value=float(p.costo_unitario))
        c_costo.number_format = fmt_moneda
        c_costo.alignment = align_right

        c_precio = ws4.cell(row=row_i, column=7, value=float(p.precio_kilo))
        c_precio.number_format = fmt_moneda
        c_precio.alignment = align_right

        ws4.cell(row=row_i, column=8, value=f"{margen:.1f}%").alignment = align_center

        c_subtot = ws4.cell(row=row_i, column=9, value=float(subtotal_stock))
        c_subtot.number_format = fmt_moneda
        c_subtot.alignment = align_right

        for col in range(1, 10):
            ws4.cell(row=row_i, column=col).border = thin_border
        row_i += 1

    ws4.cell(row=row_i, column=2, value="VALORACIÓN TOTAL ESTIMADA EN BODEGA").font = font_bold
    c_tot_inv = ws4.cell(row=row_i, column=9, value=float(total_val_bodega))
    c_tot_inv.font = font_bold
    c_tot_inv.number_format = fmt_moneda
    c_tot_inv.border = double_bottom_border

    # Ajustar anchos automáticos de columnas en todas las hojas
    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Exportar a BytesIO en memoria
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    nombre_archivo = f"Reporte_{est.nombre.replace(' ', '_')}_{d_desde.strftime('%Y%m%d')}_{d_hasta.strftime('%Y%m%d')}.xlsx"
    return output.getvalue(), nombre_archivo
