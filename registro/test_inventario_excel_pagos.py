import io
from decimal import Decimal
from django.utils import timezone
from django.test import TestCase, Client
from django.contrib.auth.models import User
from django.urls import reverse
import openpyxl

from registro.models import (
    Establecimiento,
    Municipio,
    Perfil,
    Producto,
    ItemPedido,
    Venta,
    Compra,
    TransaccionBreB,
    Auditoria,
)
from registro.excel_service import (
    generar_plantilla_pedido_excel,
    procesar_archivo_pedido_excel,
    generar_excel_conciliacion_pagos,
)


class InventarioExcelPagosTests(TestCase):
    def setUp(self):
        self.municipio = Municipio.objects.create(
            codigo_dane="15001", nombre="Tunja", departamento="Boyacá"
        )
        self.est = Establecimiento.objects.create(
            nombre="Carnicería La Esperanza",
            nit="901234567-8",
            municipio=self.municipio,
            llave_bre_b="3101234567",
            tipo_llave_bre_b="celular",
            banco_receptor_bre_b="Bancolombia",
        )

        # Propietario / Administrador
        self.user_admin = User.objects.create_user(
            username="maria.admin", password="password123", first_name="María", last_name="Gómez"
        )
        self.perfil_admin = Perfil.objects.create(
            user=self.user_admin, establecimiento=self.est, rol="propietario"
        )

        # Dependiente / Cajero
        self.user_cajero = User.objects.create_user(
            username="carlos.cajero", password="password123", first_name="Carlos", last_name="Ruiz"
        )
        self.perfil_cajero = Perfil.objects.create(
            user=self.user_cajero, establecimiento=self.est, rol="dependiente"
        )

        # Producto existente para pruebas
        self.prod_pollo = Producto.objects.create(
            establecimiento=self.est,
            nombre="Pechuga de Pollo",
            categoria="carnes",
            codigo_barras="770111222333",
            stock_kilos=Decimal("10.000"),
            precio_kilo=Decimal("18000"),
            costo_unitario=Decimal("13000"),
            unidad_medida="kg",
        )

    def test_01_descargar_plantilla_excel(self):
        """Verifica que la plantilla estándar de pedido en Excel se descargue correctamente."""
        self.client.login(username="maria.admin", password="password123")
        url = reverse("inventario_descargar_plantilla_excel")
        response = self.client.get(url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("plantilla_pedido_inventario", response["Content-Disposition"])

        # Inspeccionar el workbook en memoria
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn("Pedido e Inventario", wb.sheetnames)
        ws = wb["Pedido e Inventario"]
        self.assertEqual(ws["A4"].value, "CODIGO_BARRAS")
        self.assertEqual(ws["B4"].value, "PRODUCTO")
        self.assertGreater(ws.max_row, 4)

    def test_02_importar_pedido_excel_suma_stock_y_crea_nuevos(self):
        """Verifica que cargar un pedido Excel SUME al stock existente y cree productos nuevos."""
        # Construir un archivo Excel en memoria simulando remisión del proveedor
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["PLANTILLA DE PEDIDO"])
        ws.append(["Instrucciones"])
        ws.append([])
        ws.append(["CODIGO_BARRAS", "PRODUCTO", "CATEGORIA", "CANTIDAD_PEDIDO", "UNIDAD_MEDIDA", "COSTO_UNITARIO", "PRECIO_VENTA", "ES_SERVICIO"])
        # 1. Producto existente: Pechuga de Pollo (+15 kg recibidos)
        ws.append(["770111222333", "Pechuga de Pollo", "carnes", 15.0, "kg", 13500, 19000, "NO"])
        # 2. Producto nuevo: Lomo Fino de Res (20 kg)
        ws.append(["770999888777", "Lomo Fino de Res", "carnes", 20.0, "kg", 25000, 34000, "NO"])
        # 3. Servicio nuevo: Servicio de Transporte / Flete
        ws.append(["SERV-FLE", "Flete y Domicilio Especial", "servicios", 0.0, "servicio", 0, 8000, "SI"])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        # Item pendiente en lista de compras que debe marcarse como comprado
        item_pedido = ItemPedido.objects.create(
            establecimiento=self.est,
            producto=self.prod_pollo,
            nombre_producto="Pechuga de Pollo",
            categoria="carnes",
            cantidad_sugerida=Decimal("15"),
            estado="pendiente",
        )

        res = procesar_archivo_pedido_excel(
            establecimiento=self.est,
            archivo_bytes_o_file=buf,
            usuario=self.user_admin,
            registrar_compra=True,
            proveedor="Frigorífico Central Boyacá",
        )

        self.assertEqual(res["filas_procesadas"], 3)
        self.assertEqual(res["actualizados"], 1)
        self.assertEqual(res["creados"], 2)

        # 1. Verificar Pechuga: 10 iniciales + 15 del pedido = 25 Kilos
        self.prod_pollo.refresh_from_db()
        self.assertEqual(self.prod_pollo.stock_kilos, Decimal("25.000"))
        self.assertEqual(self.prod_pollo.precio_kilo, Decimal("19000"))

        # 2. Verificar Lomo Fino creado
        lomo = Producto.objects.get(establecimiento=self.est, codigo_barras="770999888777")
        self.assertEqual(lomo.nombre, "Lomo Fino de Res")
        self.assertEqual(lomo.stock_kilos, Decimal("20.000"))
        self.assertEqual(lomo.precio_kilo, Decimal("34000"))

        # 3. Verificar Servicio creado
        flete = Producto.objects.get(establecimiento=self.est, codigo_barras="SERV-FLE")
        self.assertTrue(flete.es_servicio)
        self.assertEqual(flete.precio_kilo, Decimal("8000"))

        # 4. Verificar que ItemPedido cambió a comprado
        item_pedido.refresh_from_db()
        self.assertEqual(item_pedido.estado, "comprado")

        # 5. Verificar que se creó el registro de Compra consolidada
        compra = Compra.objects.filter(establecimiento=self.est, proveedor="Frigorífico Central Boyacá").first()
        self.assertIsNotNone(compra)
        # Costo total: (15 * 13500) + (20 * 25000) = 202500 + 500000 = 702500
        self.assertEqual(compra.valor, Decimal("702500"))

    def test_03_permisos_segregados_administrador_vs_cajero(self):
        """
        Verifica la regla estricta:
        - Solo Administrador puede eliminar y modificar precios.
        - Dependiente/Cajero solo puede SUMAR stock (+), no puede modificar precio ni restar ni eliminar.
        """
        # A) Cajero intenta eliminar producto -> DENEGADO
        self.client.login(username="carlos.cajero", password="password123")
        url_eliminar = reverse("inventario_eliminar", kwargs={"pk": self.prod_pollo.pk})
        res_del = self.client.post(url_eliminar, follow=True)
        self.assertEqual(res_del.status_code, 200)
        self.assertContains(res_del, "Acceso denegado")
        self.assertTrue(Producto.objects.filter(pk=self.prod_pollo.pk).exists())

        # B) Cajero intenta ajustar precio -> DENEGADO
        url_ajuste = reverse("inventario_ajustar", kwargs={"pk": self.prod_pollo.pk})
        res_ajuste = self.client.post(url_ajuste, {"stock_kilos": "5", "precio_kilo": "10000"}, follow=True)
        self.assertEqual(res_ajuste.status_code, 200)
        self.assertContains(res_ajuste, "Acceso restringido")
        self.prod_pollo.refresh_from_db()
        self.assertEqual(self.prod_pollo.precio_kilo, Decimal("18000"))  # No cambió
        self.assertEqual(self.prod_pollo.stock_kilos, Decimal("10.000")) # No cambió

        # C) Cajero ingresa entrada de mercancía (+) -> PERMITIDO (Suma al stock)
        url_entrada = reverse("inventario_entrada_stock")
        res_ent = self.client.post(
            url_entrada,
            {"producto_id": self.prod_pollo.pk, "cantidad": "12.5", "nota_remision": "Remisión 481"},
            follow=True,
        )
        self.assertEqual(res_ent.status_code, 200)
        self.assertContains(res_ent, "Entrada de mercancía registrada")
        self.prod_pollo.refresh_from_db()
        self.assertEqual(self.prod_pollo.stock_kilos, Decimal("22.500"))  # 10 + 12.5

        # D) Administrador ajusta existencias y precio -> PERMITIDO
        self.client.login(username="maria.admin", password="password123")
        res_admin_ajuste = self.client.post(
            url_ajuste,
            {"stock_kilos": "30.0", "precio_kilo": "19500", "codigo_barras": "770111222333"},
            follow=True,
        )
        self.assertEqual(res_admin_ajuste.status_code, 200)
        self.prod_pollo.refresh_from_db()
        self.assertEqual(self.prod_pollo.stock_kilos, Decimal("30.000"))
        self.assertEqual(self.prod_pollo.precio_kilo, Decimal("19500"))

        # E) Administrador elimina producto -> PERMITIDO
        res_admin_del = self.client.post(url_eliminar, follow=True)
        self.assertEqual(res_admin_del.status_code, 200)
        self.assertContains(res_admin_del, "eliminado correctamente")
        self.assertFalse(Producto.objects.filter(pk=self.prod_pollo.pk).exists())

    def test_04_codigo_barras_etiqueta_y_qr_producto(self):
        """Verifica la generación de etiquetas de códigos de barras (PDF) y tarjetas QR Bre-B (PNG)."""
        self.client.login(username="maria.admin", password="password123")

        # 1. API Búsqueda por código de barras
        url_api = f"{reverse('api_buscar_producto_codigo')}?codigo=770111222333"
        res_api = self.client.get(url_api)
        self.assertEqual(res_api.status_code, 200)
        data = res_api.json()
        self.assertTrue(data["encontrado"])
        self.assertEqual(data["nombre"], "Pechuga de Pollo")
        self.assertEqual(data["precio"], 18000.0)

        # 2. Generación de tarjeta QR Bre-B para el producto
        url_qr = reverse("inventario_producto_qr", kwargs={"pk": self.prod_pollo.pk})
        res_qr = self.client.get(url_qr)
        self.assertEqual(res_qr.status_code, 200)
        self.assertEqual(res_qr["Content-Type"], "image/png")
        self.assertGreater(len(res_qr.content), 500)

        # 3. Generación de etiqueta térmica de código de barras en PDF
        url_etiqueta = reverse("inventario_producto_etiqueta_barras", kwargs={"pk": self.prod_pollo.pk})
        res_etiqueta = self.client.get(url_etiqueta)
        self.assertEqual(res_etiqueta.status_code, 200)
        self.assertEqual(res_etiqueta["Content-Type"], "application/pdf")
        self.assertGreater(len(res_etiqueta.content), 500)

    def test_05_auditoria_y_comprobacion_pagos_electronicos(self):
        """Verifica la página de auditoría de pagos electrónicos, comprobación BanRep y conciliación."""
        self.client.login(username="maria.admin", password="password123")

        # Crear venta con Bre-B y transacción asociada
        v_breb = Venta.objects.create(
            establecimiento=self.est,
            usuario=self.user_admin,
            fecha=timezone.localdate(),
            valor=Decimal("36000"),
            concepto="Venta 2 kg Pechuga Bre-B",
            medio_pago="bre_b",
        )
        tx_breb = TransaccionBreB.objects.create(
            establecimiento=self.est,
            referencia_unica="DC-2026-TEST-9988",
            token_visual_corto="#819",
            monto=Decimal("36000"),
            llave_utilizada="3101234567",
            id_transaccion_banrep="BANREP-SW-OK-889922",
            banco_origen="Nequi Bancolombia",
            payload_emvco="000201...",
            venta=v_breb,
            estado="aprobada",
        )

        # Crear venta con Nequi pendiente de conciliación
        v_nequi = Venta.objects.create(
            establecimiento=self.est,
            usuario=self.user_cajero,
            fecha=timezone.localdate(),
            valor=Decimal("50000"),
            concepto="Venta Nequi mostrador",
            medio_pago="nequi",
            conciliado_banco=False,
        )

        # 1. Acceder a la página de Auditoría de Pagos Electrónicos
        url_audit = reverse("pagos_electronicos")
        res_audit = self.client.get(url_audit)
        self.assertEqual(res_audit.status_code, 200)
        self.assertContains(res_audit, "DC-2026-TEST-9988")
        self.assertContains(res_audit, "#819")
        self.assertContains(res_audit, "$36.000")
        self.assertContains(res_audit, "BANREP-SW-OK-889922")

        # 2. Endpoint de Comprobación Criptográfica Bre-B
        url_comprobar = reverse("api_comprobar_pago_electronico", kwargs={"referencia": "DC-2026-TEST-9988"})
        res_comp = self.client.get(url_comprobar)
        self.assertEqual(res_comp.status_code, 200)
        data_comp = res_comp.json()
        self.assertTrue(data_comp["comprobado"])
        self.assertEqual(data_comp["id_transaccion_banrep"], "BANREP-SW-OK-889922")
        self.assertEqual(data_comp["protocolo_seguridad"]["algoritmo_firma"], "HMAC-SHA256")
        self.assertTrue(data_comp["protocolo_seguridad"]["firma_valida"])

        # 3. Conciliar la venta Nequi con el comprobante de extracto
        url_conciliar = reverse("pagos_electronicos_conciliar", kwargs={"pk": v_nequi.pk})
        res_conc = self.client.post(url_conciliar, {"comprobante_bancario": "APROB-M448899"}, follow=True)
        self.assertEqual(res_conc.status_code, 200)
        v_nequi.refresh_from_db()
        self.assertTrue(v_nequi.conciliado_banco)
        self.assertEqual(v_nequi.comprobante_bancario, "APROB-M448899")

        # 4. Exportar libro de conciliación bancaria a Excel
        url_excel = reverse("pagos_electronicos_exportar_excel")
        res_excel = self.client.get(url_excel)
        self.assertEqual(res_excel.status_code, 200)
        self.assertEqual(
            res_excel["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("conciliacion_pagos", res_excel["Content-Disposition"])
