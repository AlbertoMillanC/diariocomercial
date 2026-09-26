"""
Comando de Django para ejecutar el Bot de Telegram de DiarioComercial.

Comandos disponibles en el chat:
  - Venta <monto> <concepto>
  - Venta empresa <monto> <concepto> (calcula y genera retención automática)
  - Compra <monto> <proveedor>
  - Retencion <monto> <tercero>
  - /hoy (cierre de caja del día)
  - /resumen (acumulado del mes y balance)
  - /ultimas (últimos 5 movimientos con ID)
  - /ica (desglose de impuestos ICA de Tunja)
  - /anular <id_venta> <motivo> (anulación con trazabilidad de auditoría)
  - /consolidado o /excel (genera y envía el archivo Excel oficial para el contador)
"""
import io
import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request
from datetime import date
from decimal import Decimal

from django.conf import settings
from django.contrib.auth.models import User
from django.core.management.base import BaseCommand
from django.utils import timezone
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from registro.models import (
    ActividadCIIU, Auditoria, Compra, Establecimiento, Producto, Retencion, Venta,
    VinculoCanal, Perfil
)
from registro.bot_service import despachar_mensaje as bot_service_despachar
from registro.inventario_service import (
    buscar_producto_en_texto,
    consultar_producto_o_categoria,
    generar_lista_categoria,
    generar_lista_compras,
    generar_resumen_general_categorias,
    normalizar_texto,
    parsear_dinero,
    parsear_peso,
    procesar_entrada_inventario,
    procesar_salida_inventario,
    revertir_salida_inventario,
)


class TelegramClient:
    def __init__(self, token):
        self.base_url = f"https://api.telegram.org/bot{token}"

    def post_json(self, method, data):
        url = f"{self.base_url}/{method}"
        encoded_data = json.dumps(data).encode("utf-8")
        req = urllib.request.Request(
            url,
            data=encoded_data,
            headers={"Content-Type": "application/json"},
        )
        try:
            with urllib.request.urlopen(req, timeout=35) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except Exception as e:
            return {"ok": False, "error": str(e)}

    def send_message(self, chat_id, text):
        return self.post_json("sendMessage", {"chat_id": chat_id, "text": text, "parse_mode": "Markdown"})

    def send_document(self, chat_id, file_bytes, filename, caption=""):
        boundary = "----TelegramFormBoundary" + hex(int(time.time() * 1000))[2:]
        parts = []

        parts.append(
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"chat_id\"\r\n\r\n{chat_id}\r\n".encode("utf-8")
        )
        if caption:
            parts.append(
                f"--{boundary}\r\nContent-Disposition: form-data; name=\"caption\"\r\n\r\n{caption}\r\n".encode("utf-8")
            )

        mime_type = "application/pdf" if filename.lower().endswith(".pdf") else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        header = (
            f"--{boundary}\r\n"
            f"Content-Disposition: form-data; name=\"document\"; filename=\"{filename}\"\r\n"
            f"Content-Type: {mime_type}\r\n\r\n"
        ).encode("utf-8")
        parts.append(header)
        parts.append(file_bytes)
        parts.append(b"\r\n")
        parts.append(f"--{boundary}--\r\n".encode("utf-8"))

        body = b"".join(parts)
        req = urllib.request.Request(
            f"{self.base_url}/sendDocument",
            data=body,
            headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        )
        try:
            with urllib.request.urlopen(req, timeout=45) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except Exception as e:
            return {"ok": False, "error": str(e)}

    def send_photo(self, chat_id, photo_bytes, caption=""):
        boundary = "----TelegramFormBoundary" + hex(int(time.time() * 1000))[2:]
        parts = []

        parts.append(
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"chat_id\"\r\n\r\n{chat_id}\r\n".encode("utf-8")
        )
        if caption:
            parts.append(
                f"--{boundary}\r\nContent-Disposition: form-data; name=\"caption\"\r\n\r\n{caption}\r\n".encode("utf-8")
            )

        header = (
            f"--{boundary}\r\n"
            f"Content-Disposition: form-data; name=\"photo\"; filename=\"qr_bre_b.png\"\r\n"
            f"Content-Type: image/png\r\n\r\n"
        ).encode("utf-8")
        parts.append(header)
        parts.append(photo_bytes)
        parts.append(b"\r\n")
        parts.append(f"--{boundary}--\r\n".encode("utf-8"))

        body = b"".join(parts)
        req = urllib.request.Request(
            f"{self.base_url}/sendPhoto",
            data=body,
            headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        )
        try:
            with urllib.request.urlopen(req, timeout=45) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except Exception as e:
            return {"ok": False, "error": str(e)}

    def get_updates(self, offset=None, timeout=25):
        data = {"timeout": timeout}
        if offset is not None:
            data["offset"] = offset
        return self.post_json("getUpdates", data)


class Command(BaseCommand):
    help = "Inicia el bot de Telegram de DiarioComercial con soporte completo de comandos y exportación Excel"

    def add_arguments(self, parser):
        parser.add_argument(
            "--token",
            type=str,
            help="Token secreto del bot entregado por @BotFather",
        )

    def handle(self, *args, **options):
        token = options.get("token") or os.environ.get("TELEGRAM_BOT_TOKEN")
        if not token and os.path.exists(".env"):
            try:
                with open(".env", "r", encoding="utf-8") as f:
                    for line in f:
                        if line.startswith("TELEGRAM_BOT_TOKEN="):
                            token = line.split("=", 1)[1].strip().strip('"').strip("'")
                            break
            except Exception:
                pass

        if not token:
            self.stdout.write(self.style.ERROR("No se encontró ningún token en .env ni como argumento."))
            return

        client = TelegramClient(token)

        self.stdout.write(self.style.SUCCESS("=" * 65))
        self.stdout.write(self.style.SUCCESS("  BOT TELEGRAM DIARIOCOMERCIAL AVANZADO ACTIVO"))
        self.stdout.write(self.style.SUCCESS("  Escuchando ventas, compras, retenciones y /consolidado..."))
        self.stdout.write(self.style.SUCCESS("  (Presiona Ctrl + C para detener)"))
        self.stdout.write(self.style.SUCCESS("=" * 65 + "\n"))

        offset = None

        while True:
            try:
                updates = client.get_updates(offset=offset)
                if not updates.get("ok"):
                    time.sleep(3)
                    continue

                for u in updates.get("result", []):
                    offset = u["update_id"] + 1
                    msg = u.get("message") or u.get("edited_message")
                    if not msg or "text" not in msg:
                        continue

                    chat_id = msg["chat"]["id"]
                    texto = msg["text"].strip()
                    autor = msg.get("from", {}).get("first_name", "Usuario")

                    msg_id = msg.get("message_id")
                    self.stdout.write(f"[{timezone.now().strftime('%H:%M:%S')}] {autor}: '{texto}'")
                    self.despachar_mensaje(client, chat_id, texto, autor, msg_id=msg_id)

            except KeyboardInterrupt:
                self.stdout.write(self.style.WARNING("\nDeteniendo bot de Telegram..."))
                break
            except Exception as e:
                self.stdout.write(self.style.ERROR(f"Error en polling: {e}"))
                time.sleep(2)

    def despachar_mensaje(self, client, chat_id, texto, autor, msg_id=None):
        texto_limpio = texto.strip()
        cmd = texto_limpio.lower()

        # Resolver el vínculo multi-tenant del canal
        vinculo = VinculoCanal.objects.select_related("usuario", "establecimiento").filter(
            canal="telegram", identificador_externo=str(chat_id), activo=True
        ).first()
        est = vinculo.establecimiento if vinculo else None

        # Si el usuario no está vinculado, verificar si es comando de autenticación
        es_auth = bool(re.search(r"auth_[A-Za-z0-9_-]+", texto_limpio) or re.search(r"\b\d{6}\b", texto_limpio) or cmd.startswith("/start"))

        if not est and not es_auth:
            base_url = getattr(settings, "BASE_URL", os.environ.get("BASE_URL", "http://127.0.0.1:8001")).rstrip("/")
            link_activar = f"{base_url}/configuracion/"
            link_web = f"{base_url}/"
            client.send_message(
                chat_id,
                "⚠️ *Tu cuenta de Telegram aún no está vinculada a ningún comercio.*\n\n"
                "👉 *Para activarla y conectarla en 1 solo clic:*\n"
                f"1. Abre directamente este enlace: {link_activar}\n"
                "2. Inicia sesión en tu cuenta de comercio.\n"
                "3. En la pestaña *Asistente Móvil*, toca el botón *\"Conectar con Telegram Ahora\"* o escribe aquí el código PIN de 6 dígitos que te muestra en pantalla.\n\n"
                f"🌐 *¿Aún no tienes cuenta?* Conoce la plataforma y regístrate en: {link_web}\n\n"
                "_(Si ya tienes tu código PIN de 6 dígitos, digítalo y envíalo directamente en este chat)_."
            )
            return

        # 1. Comando /consolidado o /excel -> Envía archivo Excel adjunto (RBAC: Solo Propietario)
        if cmd in ("/consolidado", "consolidado", "/excel", "excel"):
            if not est:
                client.send_message(chat_id, "⚠️ Debes vincular tu comercio antes de solicitar reportes.")
                return

            perfil = Perfil.objects.filter(user=vinculo.usuario, establecimiento=est).first()
            if perfil and not perfil.es_propietario():
                client.send_message(
                    chat_id,
                    "⛔ *Acceso Restringido:* El consolidado contable en Excel está reservado exclusivamente al propietario."
                )
                return

            client.send_message(chat_id, "⏳ _Generando archivo Excel del consolidado para el contador..._")
            archivo_bytes, nombre_archivo = self.generar_excel_consolidado(est=est)
            caption = f"📊 *Consolidado Oficial de DiarioComercial*\nEstablecimiento: {est.nombre}\nPeriodo: {date.today().strftime('%B %Y')}"
            client.send_document(chat_id, archivo_bytes, nombre_archivo, caption)
            return

        # 2. Comandos de inventario por grupos / departamentos
        if cmd in ("/inventario", "inventario", "/stock", "stock"):
            client.send_message(chat_id, generar_resumen_general_categorias(est))
            return

        if cmd in ("/carnes", "carnes", "/carne", "carne"):
            client.send_message(chat_id, generar_lista_categoria(est, "carnes"))
            return

        if cmd in ("/abarrotes", "abarrotes", "/viveres", "viveres", "/granos", "granos"):
            client.send_message(chat_id, generar_lista_categoria(est, "abarrotes"))
            return

        if cmd in ("/lacteos", "lacteos", "/huevos", "huevos"):
            client.send_message(chat_id, generar_lista_categoria(est, "lacteos"))
            return

        if cmd in ("/fruver", "fruver", "/verduras", "verduras", "/frutas", "frutas"):
            client.send_message(chat_id, generar_lista_categoria(est, "fruver"))
            return

        if cmd in ("/bebidas", "bebidas", "/bebida", "bebida"):
            client.send_message(chat_id, generar_lista_categoria(est, "bebidas"))
            return

        if cmd in ("/aseo", "aseo", "/limpieza", "limpieza"):
            client.send_message(chat_id, generar_lista_categoria(est, "aseo"))
            return

        # 3. Comandos de Compras / Pedido Formal y preguntas de faltantes
        t_compras = normalizar_texto(texto_limpio).replace("¿", "").replace("?", "").strip()
        es_consulta_compras = (
            cmd in ("/comprar", "comprar", "/pedido", "pedido", "/pedidos", "pedidos", "/faltantes", "faltantes")
            or t_compras in (
                "que debo comprar", "que hay que comprar", "que se acabo", "que falta",
                "que falta comprar", "lista de compras", "lista compras", "que compras faltan", "que comprar"
            )
            or bool(re.match(r"^(?:que\s+(?:debo|hay\s+que|falta|toca)\s+comprar|que\s+se\s+acabo|lista\s+(?:de\s+)?compras?|pedidos?)$", t_compras))
        )
        if es_consulta_compras:
            client.send_message(chat_id, generar_lista_compras(est))
            return

        # 4. Consultas en lenguaje natural de inventario
        if est:
            resp_consulta = consultar_producto_o_categoria(est, texto_limpio)
            if resp_consulta:
                client.send_message(chat_id, resp_consulta)
                return

        # 5. Comando /ultimas -> Últimos 5 movimientos
        if cmd in ("/ultimas", "ultimas", "/historial", "historial"):
            client.send_message(chat_id, self.generar_ultimas(est=est))
            return

        # 6. Comando /ica -> Información de impuesto ICA
        if cmd in ("/ica", "ica", "/impuesto", "impuesto"):
            client.send_message(chat_id, self.generar_informacion_ica(est=est))
            return

        # 7. Comando /resumen -> Acumulado del mes
        if cmd in ("/resumen", "resumen", "/saldo", "saldo"):
            client.send_message(chat_id, self.generar_resumen_mes(est=est))
            return

        # 8. Comando /anular <id> <motivo>
        match_anular = re.match(r"^/anular\s+(\d+)\s*(.*)$", texto_limpio, re.IGNORECASE)
        if match_anular:
            venta_id = int(match_anular.group(1))
            motivo = match_anular.group(2).strip() or "Error de digitación en caja"
            client.send_message(chat_id, self.anular_venta(venta_id, motivo, autor, est=est))
            return

        # 9. Para todos los demás comandos (ventas, compras, stock, cierre /hoy, cobro Bre-B), delega al motor central
        resultado = bot_service_despachar(
            canal="telegram",
            identificador_externo=str(chat_id),
            texto_mensaje=texto_limpio,
            identificador_mensaje=str(msg_id) if msg_id else None,
            nombre_remitente=autor,
            return_adjuntos=True,
        )
        if isinstance(resultado, tuple):
            respuesta, venta_obj, cobro_info = resultado
        else:
            respuesta, venta_obj, cobro_info = resultado, None, None

        if respuesta:
            client.send_message(chat_id, respuesta)

        # Si se registró una venta, enviar automáticamente el QR Bre-B y la factura / recibo en PDF
        if venta_obj:
            try:
                from registro.recibo_service import preparar_paquete_omnicanal_venta
                paquete = preparar_paquete_omnicanal_venta(venta_obj)
                # 1. Enviar Código QR Bre-B (Estilo WeChat Pay para pagar en mostrador)
                client.send_photo(chat_id, paquete["qr_bytes"], caption=paquete["qr_caption"])
                # 2. Enviar Factura / Recibo de Pago en PDF
                client.send_document(chat_id, paquete["pdf_bytes"], paquete["pdf_filename"], caption=paquete["pdf_caption"])
            except Exception as e_adj:
                self.stdout.write(self.style.WARNING(f"Aviso generando adjuntos venta #{venta_obj.pk}: {e_adj}"))

        # Si fue una solicitud de cobro Bre-B directo (/cobrar o /qr)
        elif cobro_info:
            try:
                from registro.recibo_service import generar_imagen_qr_bre_b
                tx = cobro_info["tx"]
                est_cobro = cobro_info.get("establecimiento") or est
                qr_bytes = generar_imagen_qr_bre_b(
                    monto=cobro_info["monto"],
                    establecimiento=est_cobro,
                    referencia=tx.referencia_unica,
                    payload_emvco=cobro_info["payload"],
                )
                monto_fmt = f"${cobro_info['monto']:,.0f} COP".replace(",", ".")
                caption = (
                    f"⚡ *Código QR Bre-B para Cobro en Mostrador (WeChat Pay)*\n"
                    f"💰 *Monto exacto:* {monto_fmt}\n"
                    f"Ref: {tx.token_visual_corto} • {tx.referencia_unica[:18]}\n"
                    f"📲 El cliente puede escanear con Nequi, Daviplata o cualquier banco."
                )
                client.send_photo(chat_id, qr_bytes, caption=caption)
            except Exception as e_cobro:
                self.stdout.write(self.style.WARNING(f"Aviso generando QR cobro: {e_cobro}"))

        return

    def mensaje_ayuda(self, autor):
        return (
            f"👋 *¡Hola {autor}! Asistente DiarioComercial*\n"
            "_Comandos y consultas en lenguaje natural en tiempo real:_\n\n"
            "🏪 *CONSULTAS DE INVENTARIO Y SUPERMERCADO:*\n"
            "• `/inventario` ➡️ Resumen general por departamentos (kilos y valor)\n"
            "• `/carnes` ➡️ Carnes y embutidos (kilo, libra y gramo)\n"
            "• `/abarrotes` ➡️ Víveres, arroz, granos, aceites y harinas\n"
            "• `/lacteos` ➡️ Lácteos, quesos y huevos campesinos\n"
            "• `/fruver` ➡️ Frutas, verduras, papas y plátanos\n"
            "• `/bebidas` ➡️ Bebidas, gaseosas y jugos\n"
            "• `/aseo` ➡️ Productos de aseo y limpieza del hogar\n\n"
            "💬 *PREGUNTAS EN LENGUAJE NATURAL:*\n"
            "• `¿tenemos carne?` o `que carnes tengo` ➡️ Muestra cortes disponibles\n"
            "• `¿tenemos tocino?` o `tenemos arroz` ➡️ Existencias y precios exactos\n"
            "• `¿que abarrotes hay?` o `que lacteos hay` ➡️ Catálogo del grupo\n"
            "🛒 *LISTA DE COMPRAS Y PEDIDOS A PROVEEDORES:*\n"
            "• `/comprar` o `/pedidos` ➡️ Ver lista formal de compras y faltantes\n"
            "• `¿qué debo comprar?` o `¿qué se acabó?` ➡️ Reporte de compras urgentes\n"
            "• Si escribes ej: `salchicha` o `¿hay salchichas?`, se agrega automáticamente a compras futuras\n\n"
            "📝 *VENTAS RÁPIDAS (Calcula gramos y descuenta stock):*\n"
            "• `40 mil carne molida` o `40000 carne molida`\n"
            "• `20 mil tocino` o `15 mil costilla`\n"
            "• `10 mil arroz` o `1 libra de frijol`\n"
            "• `Venta 50000 viveres mostrador`\n"
            "• `Venta empresa 300000 Boyaca SAS`\n\n"
            "📊 *REPORTES CONTABLES:*\n"
            "• `/hoy` ➡️ Cierre de caja en vivo\n"
            "• `/resumen` ➡️ Balance del mes\n"
            "• `/ultimas` ➡️ Últimos movimientos con ID\n"
            "• `/ica` ➡️ Estimación del impuesto ICA Tunja\n"
            "• `/consolidado` ➡️ *Te envía el archivo Excel oficial para el contador*\n\n"
            "🚫 *ANULACIONES:*\n"
            "• `/anular <ID> <motivo>` ➡️ Anula una venta, restaura existencias y guarda auditoría"
        )

    def extraer_datos_venta(self, cuerpo):
        """
        Interpreta conceptos de venta en lenguaje natural colombiano:
        - "40 mil carne molida nequi" -> valor=40000, concepto="carne molida", medio="nequi"
        - "40000 carne molida" -> valor=40000, concepto="carne molida", medio="efectivo"
        - "20 mil tocino daviplata" -> valor=20000, concepto="tocino", medio="daviplata"
        - "1 libra carne molida" -> valor=None, concepto="1 libra carne molida"
        - "45000 viveres mostrador transferencia" -> valor=45000, concepto="viveres mostrador"
        """
        # 1. Detectar medio de pago (Nequi, Daviplata, Transferencia, Efectivo)
        cuerpo_lower = cuerpo.lower()
        medio_pago = "efectivo"
        if re.search(r"\bnequi\b", cuerpo_lower):
            medio_pago = "nequi"
        elif re.search(r"\b(?:daviplata|davi)\b", cuerpo_lower):
            medio_pago = "daviplata"
        elif re.search(r"\b(?:transferencia|bancolombia|transfiya)\b", cuerpo_lower):
            medio_pago = "transferencia"
        elif re.search(r"\befectivo\b", cuerpo_lower):
            medio_pago = "efectivo"

        # Remover palabra de medio de pago para no ensuciar el concepto
        t_sin_medio = re.sub(
            r"\b(nequi|daviplata|davi|transferencia|bancolombia|transfiya|efectivo)\b",
            "",
            cuerpo,
            flags=re.IGNORECASE,
        ).strip()

        kilos = parsear_peso(t_sin_medio)
        t_sin_peso = t_sin_medio
        if kilos is not None:
            t_sin_peso = re.sub(
                r"(una|un|dos|tres|cuatro|cinco|media|1/2|cuarto|1/4|3/4|\d+(?:[\.,]\d+)?)\s*(libras?|lb|kilos?|kg|gramos?|g)\b",
                "",
                t_sin_peso,
                flags=re.IGNORECASE,
            ).strip()

        valor, texto_sin_dinero = parsear_dinero(t_sin_peso)
        concepto = texto_sin_dinero if valor else t_sin_medio
        if not concepto or concepto.lower() in ("de", "para", "en", "por"):
            concepto = "Venta mostrador"

        return kilos, valor, concepto, medio_pago

    def guardar_venta(self, valor, concepto, autor, medio_pago="efectivo", est=None, usuario=None):
        if not est:
            return "⚠️ Establecimiento no especificado o canal no vinculado."
        usuario = usuario or User.objects.filter(perfil__establecimiento=est).first() or User.objects.first()
        actividad = ActividadCIIU.objects.filter(establecimiento=est).first()
        tarifa_mil = actividad.tarifa_x_mil if actividad else Decimal("6.0")

        # Descontar stock y autocalcular valor si se vendió por peso sin dinero explícito
        prod, kilos, libras, val_calc, info_stock = procesar_salida_inventario(est, concepto, valor)
        if (not valor or valor <= Decimal("1")) and val_calc > 0:
            valor = val_calc

        ica = (valor * tarifa_mil) / Decimal("1000")

        venta = Venta.objects.create(
            establecimiento=est,
            usuario=usuario,
            actividad=actividad,
            fecha=date.today(),
            fecha_hora=timezone.now(),
            valor=valor,
            concepto=f"{concepto} (Telegram: {autor})",
            tipo_cliente="particular",
            medio_pago=medio_pago,
            ica_estimado=ica,
            estado="vigente",
        )

        Auditoria.objects.create(
            usuario=usuario,
            entidad_afectada="venta",
            id_registro=venta.pk,
            accion="crear_telegram",
            valor_nuevo=f"valor={valor}|concepto={concepto}|medio={medio_pago}|ica={ica}|prod={prod.nombre if prod else 'none'}|kilos={kilos}",
            motivo=f"Venta registrada desde Telegram por {autor}",
        )

        icono_medio = {
            "efectivo": "💵 Efectivo",
            "nequi": "📱 Nequi",
            "daviplata": "📲 Daviplata",
            "transferencia": "🏦 Transferencia",
        }.get(medio_pago, "💵 Efectivo")

        msg = (
            "✅ *¡Venta registrada con éxito!*\n\n"
            f"🧾 *Comprobante:* #{venta.pk}\n"
            f"💵 *Valor:* ${valor:,.0f} COP\n"
            f"🏷️ *Concepto:* {concepto}\n"
            f"💳 *Medio de pago:* {icono_medio}\n"
            f"📊 *ICA estimado:* ${ica:,.2f} COP ({tarifa_mil} x mil)\n"
            f"📅 *Fecha:* {date.today().strftime('%d/%m/%Y')}"
        )
        if info_stock:
            msg += f"\n{info_stock}"
        return msg

    def guardar_venta_empresa(self, valor, concepto, autor, medio_pago="transferencia", est=None, usuario=None):
        if not est:
            return "⚠️ Establecimiento no especificado o canal no vinculado."
        usuario = usuario or User.objects.filter(perfil__establecimiento=est).first() or User.objects.first()
        actividad = ActividadCIIU.objects.filter(establecimiento=est).first()
        tarifa_mil = actividad.tarifa_x_mil if actividad else Decimal("6.0")

        prod, kilos, libras, val_calc, info_stock = procesar_salida_inventario(est, concepto, valor)
        if (not valor or valor <= Decimal("1")) and val_calc > 0:
            valor = val_calc

        ica = (valor * tarifa_mil) / Decimal("1000")
        reteica = ica

        venta = Venta.objects.create(
            establecimiento=est,
            usuario=usuario,
            actividad=actividad,
            fecha=date.today(),
            fecha_hora=timezone.now(),
            valor=valor,
            concepto=f"{concepto} [Venta Empresa] (Telegram: {autor})",
            tipo_cliente="empresa",
            medio_pago=medio_pago,
            ica_estimado=ica,
            estado="vigente",
        )

        retencion = Retencion.objects.create(
            establecimiento=est,
            usuario=usuario,
            venta=venta,
            fecha=date.today(),
            tipo="ica",
            valor=reteica,
            tercero=concepto,
            estado="vigente",
        )

        Auditoria.objects.create(
            usuario=usuario,
            entidad_afectada="venta",
            id_registro=venta.pk,
            accion="crear_telegram_empresa",
            valor_nuevo=f"valor={valor}|reteica={reteica}|prod={prod.nombre if prod else 'none'}|kilos={kilos}",
            motivo=f"Venta a empresa con retención registrada por {autor}",
        )

        msg = (
            "🏢 *¡Venta Corporativa a Empresa Registrada!*\n\n"
            f"🧾 *Comprobante:* #{venta.pk}\n"
            f"💵 *Valor Bruto:* ${valor:,.0f} COP\n"
            f"📉 *Retención ICA aplicada:* -${reteica:,.2f} COP (Retención #{retencion.pk})\n"
            f"💰 *Neto a recaudar:* ${(valor - reteica):,.0f} COP\n"
            f"📊 *ICA generado:* ${ica:,.2f} COP"
        )
        if info_stock:
            msg += f"\n{info_stock}"
        return msg

    def guardar_compra(self, valor, proveedor, autor, est=None, usuario=None):
        if not est:
            return "⚠️ Establecimiento no especificado o canal no vinculado."
        usuario = usuario or User.objects.filter(perfil__establecimiento=est).first() or User.objects.first()

        # Reabastecer stock si la compra incluye producto y peso (ej: 20 kilos tocino Frigorifico)
        prod_ent, kilos_ent, info_entrada = procesar_entrada_inventario(est, proveedor)

        compra = Compra.objects.create(
            establecimiento=est,
            usuario=usuario,
            fecha=date.today(),
            valor=valor,
            proveedor=f"{proveedor} (Telegram: {autor})",
            estado="vigente",
        )

        Auditoria.objects.create(
            usuario=usuario,
            entidad_afectada="compra",
            id_registro=compra.pk,
            accion="crear_telegram",
            valor_nuevo=f"valor={valor}|proveedor={proveedor}|prod={prod_ent.nombre if prod_ent else 'none'}|kilos={kilos_ent}",
            motivo=f"Compra registrada desde Telegram por {autor}",
        )

        msg = (
            "🛒 *¡Compra/Gasto registrado!*\n\n"
            f"🧾 *Comprobante:* #{compra.pk}\n"
            f"💵 *Valor:* ${valor:,.0f} COP\n"
            f"🏢 *Proveedor:* {proveedor}\n"
            f"📅 *Fecha:* {date.today().strftime('%d/%m/%Y')}"
        )
        if info_entrada:
            msg += f"\n\n{info_entrada}"
        return msg

    def guardar_retencion(self, valor, tercero, autor, est=None, usuario=None):
        if not est:
            return "⚠️ Establecimiento no especificado o canal no vinculado."
        usuario = usuario or User.objects.filter(perfil__establecimiento=est).first() or User.objects.first()

        ret = Retencion.objects.create(
            establecimiento=est,
            usuario=usuario,
            fecha=date.today(),
            tipo="ica",
            valor=valor,
            tercero=f"{tercero} (Telegram: {autor})",
            estado="vigente",
        )

        Auditoria.objects.create(
            usuario=usuario,
            entidad_afectada="retencion",
            id_registro=ret.pk,
            accion="crear_telegram",
            valor_nuevo=f"valor={valor}|tercero={tercero}",
            motivo=f"Retención registrada desde Telegram por {autor}",
        )

        return (
            "📑 *¡Retención Registrada!*\n\n"
            f"🧾 *Comprobante:* #{ret.pk}\n"
            f"💵 *Valor retenido:* ${valor:,.0f} COP\n"
            f"🏢 *Tercero:* {tercero}\n"
            "_Suma a tu favor en el consolidado del periodo._"
        )

    def anular_venta(self, venta_id, motivo, autor, est=None):
        if not est:
            return "⚠️ Establecimiento no especificado o canal no vinculado."
        usuario = User.objects.filter(perfil__establecimiento=est, perfil__rol__in=["propietario", "empresario"]).first() or User.objects.first()
        venta = Venta.objects.filter(establecimiento=est, pk=venta_id).first()

        if not venta:
            return f"❌ No se encontró la venta con ID #{venta_id}."

        if venta.estado == "anulado":
            return f"ℹ️ La venta #{venta_id} ya se encuentra anulada."

        venta.estado = "anulado"
        venta.save()

        # Si tenía retención asociada, se anula también
        Retencion.objects.filter(venta=venta, estado="vigente").update(estado="anulado")

        # Revertir existencias de inventario si aplica
        info_reversion = revertir_salida_inventario(est, venta)

        Auditoria.objects.create(
            usuario=usuario,
            entidad_afectada="venta",
            id_registro=venta.pk,
            accion="anular_telegram",
            valor_nuevo=f"estado=anulado|{info_reversion}",
            motivo=f"Anulada desde Telegram por {autor}: {motivo}",
        )

        msg = (
            f"🚫 *¡Venta #{venta_id} Anulada Exitosamente!*\n\n"
            f"💵 *Valor anulado:* ${venta.valor:,.0f} COP\n"
            f"📋 *Motivo:* {motivo}\n"
            f"🛡️ *Auditoría:* Registro histórico preservado en base de datos."
        )
        if info_reversion:
            msg += f"\n{info_reversion}"
        return msg

    def generar_resumen_hoy(self, est=None):
        if not est:
            return "⚠️ Establecimiento no especificado o canal no vinculado."
        hoy = date.today()
        ventas = Venta.objects.filter(establecimiento=est, fecha=hoy, estado="vigente")
        compras = Compra.objects.filter(establecimiento=est, fecha=hoy, estado="vigente")

        total_ventas = sum((v.valor for v in ventas), Decimal("0"))
        ica_hoy = sum((v.ica_estimado for v in ventas), Decimal("0"))
        total_compras = sum((c.valor for c in compras), Decimal("0"))
        cant = ventas.count()
        promedio = (total_ventas / cant) if cant > 0 else Decimal("0")

        total_efectivo = sum((v.valor for v in ventas if v.medio_pago == "efectivo"), Decimal("0"))
        total_nequi = sum((v.valor for v in ventas if v.medio_pago == "nequi"), Decimal("0"))
        total_daviplata = sum((v.valor for v in ventas if v.medio_pago == "daviplata"), Decimal("0"))
        total_transf = sum((v.valor for v in ventas if v.medio_pago == "transferencia"), Decimal("0"))

        desglose = (
            "💳 *ARQUEO POR MEDIO DE PAGO:*\n"
            f"   • 💵 Efectivo en cajón: *${total_efectivo:,.0f} COP*\n"
            f"   • 📱 Nequi: *${total_nequi:,.0f} COP*\n"
            f"   • 📲 Daviplata: *${total_daviplata:,.0f} COP*\n"
            f"   • 🏦 Transferencia: *${total_transf:,.0f} COP*\n\n"
        )

        return (
            f"☀️ *Cierre de Caja de Hoy ({hoy.strftime('%d/%m/%Y')})*\n"
            f"📍 _{est.nombre}_\n\n"
            f"🧾 *Operaciones:* {cant} ventas registradas\n"
            f"💵 *Total Ventas:* ${total_ventas:,.0f} COP\n"
            f"🎯 *Ticket Promedio:* ${promedio:,.0f} COP\n\n"
            f"{desglose}"
            f"🛒 *Gastos del día:* ${total_compras:,.0f} COP\n"
            f"📊 *ICA Estimado Hoy:* ${ica_hoy:,.2f} COP\n"
            f"💰 *Caja Neta Hoy:* ${(total_ventas - total_compras):,.0f} COP"
        )

    def generar_resumen_mes(self, est=None):
        if not est:
            return "⚠️ Establecimiento no especificado o canal no vinculado."
        hoy = date.today()
        inicio_mes = hoy.replace(day=1)

        ventas = Venta.objects.filter(establecimiento=est, fecha__range=(inicio_mes, hoy), estado="vigente")
        compras = Compra.objects.filter(establecimiento=est, fecha__range=(inicio_mes, hoy), estado="vigente")
        rets = Retencion.objects.filter(establecimiento=est, fecha__range=(inicio_mes, hoy), estado="vigente")

        ingresos = sum((v.valor for v in ventas), Decimal("0"))
        ica = sum((v.ica_estimado for v in ventas), Decimal("0"))
        egresos = sum((c.valor for c in compras), Decimal("0"))
        retenciones = sum((r.valor for r in rets), Decimal("0"))

        return (
            f"📅 *Consolidado del Mes ({inicio_mes.strftime('%d/%m')} - {hoy.strftime('%d/%m/%Y')})*\n"
            f"📍 _{est.nombre}_\n\n"
            f"📈 *Ingresos Brutos:* ${ingresos:,.0f} COP\n"
            f"🛒 *Compras y Gastos:* ${egresos:,.0f} COP\n"
            f"📑 *Retenciones aplicadas:* ${retenciones:,.0f} COP\n"
            f"⚖️ *Base Neta Estimada:* ${(ingresos - egresos):,.0f} COP\n"
            f"🏛️ *Impuesto ICA Estimado:* ${ica:,.2f} COP"
        )

    def generar_ultimas(self, est=None):
        if not est:
            return "⚠️ Establecimiento no especificado o canal no vinculado."
        ventas = Venta.objects.filter(establecimiento=est).order_by("-fecha_hora", "-id")[:5]

        if not ventas:
            return "📭 No hay movimientos registrados todavía."

        lineas = ["🕒 *Últimos 5 Movimientos Registrados:*\n"]
        for v in ventas:
            icono = "✅" if v.estado == "vigente" else "❌ (Anulada)"
            lineas.append(
                f"{icono} *Venta #{v.pk}* — ${v.valor:,.0f} COP\n"
                f"   _{v.concepto}_ | {v.fecha_hora.strftime('%d/%m %H:%M')}"
            )
        return "\n\n".join(lineas)

    def generar_informacion_ica(self, est=None):
        if not est:
            return "⚠️ Establecimiento no especificado o canal no vinculado."
        hoy = date.today()
        inicio_mes = hoy.replace(day=1)

        actividades = ActividadCIIU.objects.filter(establecimiento=est)
        lineas = [
            f"🏛️ *Estructura del Impuesto ICA — Tunja*\n"
            f"📍 Establecimiento: *{est.nombre}*\n"
            f"NIT: {est.nit or 'No configurado'}\n\n"
            "📋 *Actividades Económicas y Tarifas:*"
        ]

        for act in actividades:
            lineas.append(f"• *{act.codigo}*: {act.descripcion}\n  Tarifa: *{act.tarifa_x_mil} por mil*")

        ventas_mes = Venta.objects.filter(establecimiento=est, fecha__range=(inicio_mes, hoy), estado="vigente")
        ica_total = sum((v.ica_estimado for v in ventas_mes), Decimal("0"))
        lineas.append(f"\n📊 *Total ICA acumulado este mes:* ${ica_total:,.2f} COP")
        lineas.append("_Este valor es la base informativa para preparar la declaración municipal._")
        return "\n".join(lineas)

    def generar_inventario_carnes(self, est=None):
        if not est:
            return "⚠️ Establecimiento no especificado o canal no vinculado."
        productos = Producto.objects.filter(establecimiento=est, categoria="carnes", estado="activo").order_by("nombre")
        if not productos:
            return "🥩 No hay cortes de carne registrados en el inventario."

        lineas = [
            "🥩 *LISTA DE PRECIOS Y EXISTENCIAS — CARNES*",
            f"📍 _{est.nombre}_\n",
        ]
        for p in productos:
            lineas.append(
                f"• *{p.nombre}*\n"
                f"  📦 Stock: *{p.stock_kilos} Kg* ({p.stock_libras} lb)\n"
                f"  💰 *Kilo:* ${p.precio_kilo:,.0f} | *Libra:* ${p.precio_libra:,.0f} | *Gramo:* ${p.precio_gramo}"
            )
        lineas.append("\n⚖️ _Precios vigentes para pesaje en báscula de mostrador._")
        return "\n".join(lineas)

    def generar_inventario_general(self, est=None):
        if not est:
            return "⚠️ Establecimiento no especificado o canal no vinculado."
        productos = Producto.objects.filter(establecimiento=est, estado="activo").order_by("categoria", "nombre")
        if not productos:
            return "📦 No hay productos registrados en el inventario."

        total_kilos = sum((p.stock_kilos for p in productos), Decimal("0"))
        valor_total = sum((p.stock_kilos * p.precio_kilo for p in productos), Decimal("0"))

        lineas = [
            "📦 *RESUMEN GENERAL DE INVENTARIO*",
            f"📍 _{est.nombre}_\n",
            f"📊 Total Existencias: *{total_kilos:,.1f} Kg*",
            f"💵 Valor Estimado en Bodega: *${valor_total:,.0f} COP*",
            f"🏷️ Variedades activas: *{productos.count()} productos*\n",
            "💡 _Para ver cortes de carne específicos con precios por Kilo, Libra y Gramo, escribe:_ `/carnes`",
        ]
        return "\n".join(lineas)

    def generar_excel_consolidado(self, est=None):
        """Genera un archivo Excel (.xlsx) oficial con diseño profesional para el contador."""
        if not est:
            raise ValueError("Establecimiento requerido para generar reporte consolidado")
        wb = openpyxl.Workbook()
        hoy = date.today()
        inicio_mes = hoy.replace(day=1)

        # Paleta de estilos
        fill_header = PatternFill(start_color="12323A", end_color="12323A", fill_type="solid")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_titulo = Font(name="Calibri", size=14, bold=True, color="12323A")
        font_bold = Font(name="Calibri", size=11, bold=True)
        thin_border = Border(
            left=Side(style="thin", color="D9D0C3"),
            right=Side(style="thin", color="D9D0C3"),
            top=Side(style="thin", color="D9D0C3"),
            bottom=Side(style="thin", color="D9D0C3"),
        )

        # ---------------- HOJA 1: RESUMEN EJECUTIVO ----------------
        ws1 = wb.active
        ws1.title = "Resumen Periodo"
        ws1.views.sheetView[0].showGridLines = True

        ws1["A1"] = f"DIARIO COMERCIAL — REPORTE CONSOLIDADO PARA CONTADOR"
        ws1["A1"].font = font_titulo
        ws1["A2"] = f"Establecimiento: {est.nombre} | NIT: {est.nit or '—'}"
        ws1["A3"] = f"Periodo Liquidado: {inicio_mes.strftime('%d/%m/%Y')} al {hoy.strftime('%d/%m/%Y')}"

        ventas = Venta.objects.filter(establecimiento=est, fecha__range=(inicio_mes, hoy), estado="vigente")
        compras = Compra.objects.filter(establecimiento=est, fecha__range=(inicio_mes, hoy), estado="vigente")
        rets = Retencion.objects.filter(establecimiento=est, fecha__range=(inicio_mes, hoy), estado="vigente")

        ingresos = sum((v.valor for v in ventas), Decimal("0"))
        ica = sum((v.ica_estimado for v in ventas), Decimal("0"))
        egresos = sum((c.valor for c in compras), Decimal("0"))
        retenciones = sum((r.valor for r in rets), Decimal("0"))
        neto = ingresos - egresos

        headers_resumen = ["Concepto Contable", "Valor ($ COP)", "Observación Fiscal"]
        for col_idx, h in enumerate(headers_resumen, 1):
            cell = ws1.cell(row=5, column=col_idx, value=h)
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = Alignment(horizontal="center")

        datos_resumen = [
            ("Ingresos Brutos Operacionales (Ventas)", float(ingresos), "Suma de ventas vigentes"),
            ("ICA Estimado (Alcaldía de Tunja)", float(ica), "Cálculo según tarifa CIIU"),
            ("Compras y Gastos Deducibles", float(egresos), "Soportes de mercancía y costos"),
            ("Retenciones Practicadas (ReteICA)", float(retenciones), "Descuento en declaración"),
            ("Base Neta Estimada del Periodo", float(neto), "Ingresos menos egresos"),
        ]

        for row_idx, fila in enumerate(datos_resumen, 6):
            ws1.cell(row=row_idx, column=1, value=fila[0]).font = font_bold
            c2 = ws1.cell(row=row_idx, column=2, value=fila[1])
            c2.number_format = '"$"#,##0'
            ws1.cell(row=row_idx, column=3, value=fila[2])
            for col_idx in range(1, 4):
                ws1.cell(row=row_idx, column=col_idx).border = thin_border

        # Pie de página en Hoja 1
        pie_row_ws1 = len(datos_resumen) + 8
        c_pie1 = ws1.cell(row=pie_row_ws1, column=1, value="CARLOS ALBERTO MILLAN CASTAÑO / DESARROLLO WEB")
        c_pie1.font = Font(name="Calibri", size=10, bold=True, color="5C6B70")

        # ---------------- HOJA 2: DETALLE DE MOVIMIENTOS ----------------
        ws2 = wb.create_sheet(title="Detalle Movimientos")
        ws2.views.sheetView[0].showGridLines = True

        headers_det = ["Fecha", "Tipo", "ID", "Concepto / Proveedor", "Valor ($ COP)", "Estado", "Responsable"]
        for col_idx, h in enumerate(headers_det, 1):
            cell = ws2.cell(row=1, column=col_idx, value=h)
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = Alignment(horizontal="center")

        movimientos = []
        for v in Venta.objects.filter(establecimiento=est, fecha__range=(inicio_mes, hoy)):
            movimientos.append((v.fecha, "Venta", v.pk, v.concepto, float(v.valor), v.estado, v.usuario.username))
        for c in Compra.objects.filter(establecimiento=est, fecha__range=(inicio_mes, hoy)):
            movimientos.append((c.fecha, "Compra", c.pk, c.proveedor, float(c.valor), c.estado, c.usuario.username))
        for r in Retencion.objects.filter(establecimiento=est, fecha__range=(inicio_mes, hoy)):
            movimientos.append((r.fecha, "Retención", r.pk, r.tercero, float(r.valor), r.estado, r.usuario.username))

        movimientos.sort(key=lambda x: str(x[0]), reverse=True)

        for row_idx, m in enumerate(movimientos, 2):
            ws2.cell(row=row_idx, column=1, value=m[0].strftime("%d/%m/%Y"))
            ws2.cell(row=row_idx, column=2, value=m[1])
            ws2.cell(row=row_idx, column=3, value=f"#{m[2]}")
            ws2.cell(row=row_idx, column=4, value=m[3])
            c_val = ws2.cell(row=row_idx, column=5, value=m[4])
            c_val.number_format = '"$"#,##0'
            ws2.cell(row=row_idx, column=6, value=m[5])
            ws2.cell(row=row_idx, column=7, value=m[6])
            for col_idx in range(1, 8):
                ws2.cell(row=row_idx, column=col_idx).border = thin_border

        # Pie de página en Hoja 2
        pie_row_ws2 = len(movimientos) + 3
        c_pie2 = ws2.cell(row=pie_row_ws2, column=1, value="CARLOS ALBERTO MILLAN CASTAÑO / DESARROLLO WEB")
        c_pie2.font = Font(name="Calibri", size=10, bold=True, color="5C6B70")

        # Pie de página de impresión
        ws1.oddFooter.center.text = "CARLOS ALBERTO MILLAN CASTAÑO / DESARROLLO WEB"
        ws2.oddFooter.center.text = "CARLOS ALBERTO MILLAN CASTAÑO / DESARROLLO WEB"

        # Autoajuste de columnas en ambas hojas
        for ws in (ws1, ws2):
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        nombre_archivo = f"Consolidado_Contador_{est.nombre.replace(' ', '_')}_{hoy.strftime('%Y%m')}.xlsx"
        return buf.getvalue(), nombre_archivo
