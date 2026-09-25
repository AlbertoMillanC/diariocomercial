from datetime import date
from decimal import Decimal
import io
import openpyxl

from django.contrib.auth.models import User
from django.test import TestCase, Client
from django.urls import reverse
from django.utils import timezone

from registro.models import (
    Establecimiento,
    Municipio,
    Perfil,
    ActividadCIIU,
    MotivoVenta,
    Cliente,
    Venta,
)
from registro.tax_engine import (
    obtener_o_crear_consumidor_final,
    liquidar_exogena_dian_formato_1007,
    generar_excel_exogena_formato_1007,
)


class FacturacionExogenaTestCase(TestCase):
    def setUp(self):
        self.municipio = Municipio.objects.create(
            codigo_dane="15001",
            nombre="Tunja",
            departamento="Boyacá",
        )
        self.est = Establecimiento.objects.create(
            nombre="Carnicería La Floresta",
            nit="901234567",
            municipio=self.municipio,
            direccion="Carrera 10 # 20-30",
            facturacion_electronica_habilitada=True,
            prefijo_facturacion="SETT",
            rango_desde=1,
            rango_hasta=1000,
            consecutivo_actual=10,
        )
        self.actividad = ActividadCIIU.objects.create(
            establecimiento=self.est,
            codigo="4722",
            descripcion="Comercio al por menor de carnes",
            tarifa_x_mil=Decimal("5.0"),
        )
        self.motivo = MotivoVenta.objects.create(
            establecimiento=self.est,
            actividad=self.actividad,
            nombre="Venta Carne Res",
            es_predeterminado=True,
        )

        # Usuario Propietario
        self.user = User.objects.create_user(username="don_ramon", password="password123")
        self.perfil = Perfil.objects.create(
            user=self.user,
            establecimiento=self.est,
            rol="propietario",
        )
        self.client = Client()
        self.client.login(username="don_ramon", password="password123")

    def test_obtener_o_crear_consumidor_final(self):
        """Verifica que el consumidor final normativo se cree con 12 doses (222222222222)."""
        cf = obtener_o_crear_consumidor_final(self.est)
        self.assertEqual(cf.nit_cedula, "222222222222")
        self.assertEqual(cf.nombre, "CONSUMIDOR FINAL")
        self.assertTrue(cf.es_consumidor_final)

        # Al invocar de nuevo no debe duplicarse
        cf2 = obtener_o_crear_consumidor_final(self.est)
        self.assertEqual(cf.pk, cf2.pk)

    def test_crear_cliente_individual_dian(self):
        """Registro de cliente para facturación electrónica con validación DIAN."""
        cli = Cliente.objects.create(
            establecimiento=self.est,
            tipo_documento="31",
            nit_cedula="900888777",
            dv="5",
            nombre="Restaurante Boyacense S.A.S.",
            tipo_persona="juridica",
            regimen_fiscal="48",
            correo_electronico="contabilidad@restaurante.com",
            telefono="3119998877",
            direccion="Calle Real # 4-20",
            municipio_nombre="Tunja",
            departamento_nombre="Boyacá",
        )
        self.assertFalse(cli.es_consumidor_final)
        self.assertEqual(cli.nit_cedula, "900888777")
        self.assertEqual(cli.correo_electronico, "contabilidad@restaurante.com")

    def test_venta_mostrador_vs_factura_electronica(self):
        """Comprueba la emisión de venta estándar (mostrador) vs Factura Electrónica con CUFE."""
        # 1. Venta Mostrador (Consumidor Final)
        cf = obtener_o_crear_consumidor_final(self.est)
        v1 = Venta.objects.create(
            establecimiento=self.est,
            actividad=self.actividad,
            motivo=self.motivo,
            fecha_hora=timezone.now(),
            valor=Decimal("25000"),
            cliente=cf,
            solicita_factura_electronica=False,
            usuario=self.user,
        )
        self.assertFalse(v1.solicita_factura_electronica)
        self.assertFalse(v1.numero_factura_electronica)

        # 2. Venta con Factura Electrónica Individual
        cli_empresa = Cliente.objects.create(
            establecimiento=self.est,
            tipo_documento="31",
            nit_cedula="800111222",
            dv="1",
            nombre="Asados Tunja",
            correo_electronico="factura@asados.com",
        )
        consecutivo = self.est.siguiente_consecutivo_factura()
        v2 = Venta.objects.create(
            establecimiento=self.est,
            actividad=self.actividad,
            motivo=self.motivo,
            fecha_hora=timezone.now(),
            valor=Decimal("80000"),
            cliente=cli_empresa,
            solicita_factura_electronica=True,
            numero_factura_electronica=consecutivo,
            cufe=Venta.generar_cufe(self.est, consecutivo, Decimal("80000"), timezone.now(), cli_empresa.nit_cedula),
            estado_dian="emitida",
            usuario=self.user,
        )
        self.assertTrue(v2.solicita_factura_electronica)
        self.assertEqual(v2.numero_factura_electronica, "SETT-00010")
        self.assertTrue(len(v2.cufe) >= 40)
        self.assertEqual(v2.cliente.nit_cedula, "800111222")

    def test_liquidacion_exogena_1007_suma_mostrador_y_separa_clientes(self):
        """
        REGLA CLAVE DE USUARIO:
        Suma las compras de los clientes de mostrador no identificados en una sola fila (222222222),
        a diferencia de los que piden factura que sí deben ser reportados individualmente.
        """
        cf = obtener_o_crear_consumidor_final(self.est)
        ahora = timezone.now()
        año_actual = ahora.year

        # Tres ventas en mostrador a consumidor final (no identificados)
        Venta.objects.create(
            establecimiento=self.est,
            actividad=self.actividad,
            motivo=self.motivo,
            fecha_hora=ahora,
            valor=Decimal("15000"),
            cliente=cf,
            solicita_factura_electronica=False,
            usuario=self.user,
        )
        Venta.objects.create(
            establecimiento=self.est,
            actividad=self.actividad,
            motivo=self.motivo,
            fecha_hora=ahora,
            valor=Decimal("35000"),
            cliente=cf,
            solicita_factura_electronica=False,
            usuario=self.user,
        )
        Venta.objects.create(
            establecimiento=self.est,
            actividad=self.actividad,
            motivo=self.motivo,
            fecha_hora=ahora,
            valor=Decimal("50000"),
            cliente=None,  # Venta sin cliente asignado explícitamente -> mostrador
            solicita_factura_electronica=False,
            usuario=self.user,
        )
        # Total Mostrador esperado: 15.000 + 35.000 + 50.000 = 100.000

        # Dos clientes con Factura Electrónica individual
        cli1 = Cliente.objects.create(
            establecimiento=self.est,
            tipo_documento="13",
            nit_cedula="1049123456",
            nombre="Pedro Pérez",
            correo_electronico="pedro@gmail.com",
        )
        Venta.objects.create(
            establecimiento=self.est,
            actividad=self.actividad,
            motivo=self.motivo,
            fecha_hora=ahora,
            valor=Decimal("70000"),
            cliente=cli1,
            solicita_factura_electronica=True,
            numero_factura_electronica="SETT-1",
            usuario=self.user,
        )

        cli2 = Cliente.objects.create(
            establecimiento=self.est,
            tipo_documento="31",
            nit_cedula="901999888",
            dv="3",
            nombre="Hoteles Boyacá S.A.S.",
            correo_electronico="pagos@hoteles.com",
        )
        Venta.objects.create(
            establecimiento=self.est,
            actividad=self.actividad,
            motivo=self.motivo,
            fecha_hora=ahora,
            valor=Decimal("120000"),
            cliente=cli2,
            solicita_factura_electronica=True,
            numero_factura_electronica="SETT-2",
            usuario=self.user,
        )
        # Segunda factura al mismo cli2 para verificar agregación por adquirente
        Venta.objects.create(
            establecimiento=self.est,
            actividad=self.actividad,
            motivo=self.motivo,
            fecha_hora=ahora,
            valor=Decimal("30000"),
            cliente=cli2,
            solicita_factura_electronica=True,
            numero_factura_electronica="SETT-3",
            usuario=self.user,
        )
        # Total cli2 esperado: 120.000 + 30.000 = 150.000

        # Ejecutar motor tributario Formato 1007
        resultado = liquidar_exogena_dian_formato_1007(self.est, año_actual)

        # 1. Total ventas contables: 100.000 (mostrador) + 70.000 (cli1) + 150.000 (cli2) = 320.000
        self.assertEqual(resultado["total_ingresos_brutos"], Decimal("320000"))

        # 2. Total Mostrador Consumidor Final (Suma compras de clientes no identificados)
        self.assertEqual(resultado["total_mostrador_consumidor_final"], Decimal("100000"))
        self.assertEqual(resultado["conteo_ventas_mostrador"], 3)

        # 3. Fila consolidada oficial: NIT 222222222 (9 doses normativos de exógena)
        r_mostrador = resultado["registro_consolidado_mostrador"]
        self.assertEqual(r_mostrador["nit"], "222222222")
        self.assertEqual(r_mostrador["ingresos_brutos"], Decimal("100000"))
        self.assertEqual(r_mostrador["concepto"], "4001")

        # 4. Clientes individuales discriminados
        indiv = resultado["clientes_individuales"]
        self.assertEqual(len(indiv), 2)

        # Verificar cli2 (Hoteles Boyacá) acumuló ambas facturas (120k + 30k = 150k)
        cli2_res = next(c for c in indiv if c["nit"] == "901999888")
        self.assertEqual(cli2_res["total_compras"], Decimal("150000"))
        self.assertEqual(cli2_res["cantidad_facturas"], 2)

        # Verificar cli1 (Pedro Pérez: 70k)
        cli1_res = next(c for c in indiv if c["nit"] == "1049123456")
        self.assertEqual(cli1_res["total_compras"], Decimal("70000"))
        self.assertEqual(cli1_res["cantidad_facturas"], 1)

        # 5. Cuadre 100% exacto
        self.assertTrue(resultado["cuadre_exacto"])
        self.assertEqual(resultado["diferencia_cuadre"], Decimal("0"))
        self.assertEqual(resultado["total_reportado_exogena"], Decimal("320000"))

    def test_generar_excel_exogena_formato_1007(self):
        """Verifica que la exportación a Excel (.xlsx) genere las celdas y filas correctas."""
        cf = obtener_o_crear_consumidor_final(self.est)
        Venta.objects.create(
            establecimiento=self.est,
            actividad=self.actividad,
            motivo=self.motivo,
            fecha_hora=timezone.now(),
            valor=Decimal("45000"),
            cliente=cf,
            solicita_factura_electronica=False,
            usuario=self.user,
        )
        año = timezone.now().year
        excel_bytes = generar_excel_exogena_formato_1007(self.est, año)
        self.assertTrue(len(excel_bytes) > 1000)

        # Leer con openpyxl para verificar contenido
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active
        self.assertIn("INFORMACIÓN EXÓGENA FORMATO 1007", ws["A1"].value)
        # Verificar fila de mostrador (fila 5)
        self.assertEqual(ws["C5"].value, "222222222")
        self.assertEqual(ws["J5"].value, 45000.0)

    def test_vistas_clientes_y_exogena_http(self):
        """Verifica que las rutas HTTP carguen correctamente."""
        # 1. Directorio de clientes
        resp = self.client.get(reverse("clientes_lista"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Directorio de Clientes")

        # 2. Formulario nuevo cliente
        resp = self.client.get(reverse("cliente_crear"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Registrar Cliente para Facturación Electrónica")

        # 3. API creación rápida de cliente
        data_quick = {
            "tipo_documento": "13",
            "nit_cedula": "1049777888",
            "nombre": "Ana María Gómez",
            "correo_electronico": "ana@gmail.com",
            "telefono": "3201112233",
        }
        resp = self.client.post(reverse("api_cliente_crear_rapido"), data_quick)
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.json()["exito"])
        self.assertEqual(resp.json()["nit_cedula"], "1049777888")

        # 4. Consola Exógena DIAN 1007
        resp = self.client.get(reverse("exogena_dian"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Información Exógena DIAN — Formato 1007")
        self.assertContains(resp, "222222222")

        # 5. Descarga de Excel Exógena
        resp = self.client.get(reverse("exogena_dian_exportar_excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
